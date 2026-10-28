"""Rebuild the official-list extracts in sosia/dizionari/ from the original files, after checking their SHA-256.

The original files are NOT in the repository (they are large, and only short extracts are needed). Download them
yourself from the official sources below into one folder, then run:

    python scripts/build_dictionaries.py <folder with the downloaded files>

The script refuses to run if a file does not match the SHA-256 recorded in SORGENTI (the snapshot used for this
release); it then writes the extracts and sosia/dizionari/FONTI.json with the hashes of sources and products.
Attributions and licences: NOTICE.

| File | Source | Licence |
|---|---|---|
| Elenco-comuni-italiani.xlsx | https://www.istat.it/storage/codici-unita-amministrative/Elenco-comuni-italiani.xlsx | CC BY 4.0 |
| LocalitaPuntuali_21.zip | https://www.istat.it/storage/cartografia/basi_territoriali/2021/LocalitaPuntuali_21.zip | CC BY 4.0 |
| opname_csv_gb.zip | OS Open Names, CSV (https://osdatahub.os.uk/downloads/open/OpenNames) | OGL v3.0 |
| census_2010_surnames.zip | https://www2.census.gov/topics/genealogy/2010surnames/names.zip (renamed) | public domain |
| babynames1996to2025.xlsx | ONS, "Baby names in England and Wales" (https://www.ons.gov.uk/peoplepopulationandcommunity/birthsdeathsandmarriages/livebirths/datasets/babynamesinenglandandwalesfrom1996) | OGL v3.0 |
| wikidata_cognomi_it.csv, wikidata_nomi_it.csv | Wikidata SPARQL endpoint (https://query.wikidata.org), CSV with a header and one label per line | CC0 |

Wikidata queries (the data change over time, so a new run gives different files and hashes):

    # Italian surnames: family names (Q101352) whose language (P407) is Italian (Q652)
    SELECT DISTINCT ?label WHERE { ?s wdt:P31 wd:Q101352; wdt:P407 wd:Q652; rdfs:label ?label .
                                   FILTER(LANG(?label) = "it") }
    # Italian first names: given names (Q202444 and subclasses) whose language is Italian
    SELECT DISTINCT ?label WHERE { ?s wdt:P31/wdt:P279* wd:Q202444; wdt:P407 wd:Q652; rdfs:label ?label .
                                   FILTER(LANG(?label) = "it") }
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

DIZ = Path(__file__).resolve().parent.parent / "sosia" / "dizionari"

SORGENTI = {
    "Elenco-comuni-italiani.xlsx": "83842076860450f7e482daecea6b7a769f5f93d0bf5b0d48802b44896d7a26d5",
    "LocalitaPuntuali_21.zip": "67ba8a957b9c1ea21ed6144cc33ec1ac3c3a0b1370ba4443452e63ae37e2b9f9",
    "opname_csv_gb.zip": "81ee5ba2c96d8fe5ddbf83fb4f9928e8740563618c6c5eb9376280b399a74532",
    "census_2010_surnames.zip": "117c41cb4668727b7627b2845b6df3f83eb2a22a1813f42c0ff4bdcab86de135",
    "babynames1996to2025.xlsx": "bf7b10486ab20d0b32e4bc87385112f4555210f61e34eaac3574a0d8a658c0d4",
    "wikidata_cognomi_it.csv": "a0114892a299a6cad939e9a59cd9a5423b1d93465d8a9f26ed1793b1eb0fb6c0",
    "wikidata_nomi_it.csv": "8127113034d9cd733f2aa7de5ba351fcda2a4d5e1fb00fae73c0d02544a0bb50",
}
FONTI = {
    "Elenco-comuni-italiani.xlsx": "ISTAT, Elenco dei comuni italiani (21/02/2026), CC BY 4.0",
    "LocalitaPuntuali_21.zip": "ISTAT, Basi territoriali 2021, località puntuali, CC BY 4.0",
    "opname_csv_gb.zip": "OS Open Names 2026-10, Contains OS data © Crown copyright and database right 2026, OGL v3.0",
    "census_2010_surnames.zip": "U.S. Census Bureau, 2010 surnames (File B), public domain",
    "babynames1996to2025.xlsx": "ONS, Baby names in England and Wales 1996-2025, OGL v3.0",
    "wikidata_cognomi_it.csv": "Wikidata (query 28/10/2026: Q101352, P407=Q652), CC0",
    "wikidata_nomi_it.csv": "Wikidata (query 28/10/2026: Q202444 e sottoclassi, P407=Q652), CC0",
}


def sha(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for blocco in iter(lambda: f.read(1 << 20), b""):
            h.update(blocco)
    return h.hexdigest()


def scrivi(nome: str, righe) -> str:
    righe = sorted(set(righe), key=lambda r: (r.lower(), r))
    (DIZ / nome).write_text("\n".join(righe) + "\n", encoding="utf-8", newline="\r\n")  # stessi byte su ogni sistema
    print(f"{nome}: {len(righe)} righe")
    return nome


def comuni(cart: Path) -> tuple[list, dict]:
    import openpyxl
    ws = openpyxl.load_workbook(cart / "Elenco-comuni-italiani.xlsx", read_only=True).active
    righe, sigla_di = [], {}
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r[0]:
            continue
        nome_it, altro, sigla, codnum, catast = r[6], r[7], r[14], r[15], r[20]
        sigla_di[int(codnum)] = sigla
        righe.append(f"{nome_it};{sigla};{catast}")
        if altro:  # «Bozen», «Aoste»: la denominazione nell'altra lingua
            for a in re.split(r"\s*/\s*", str(altro)):
                if a and a != nome_it:
                    righe.append(f"{a};{sigla};{catast}")
    return righe, sigla_di


def localita(cart: Path, sigla_di: dict, nomi_comuni: set) -> list:
    z = zipfile.ZipFile(cart / "LocalitaPuntuali_21.zip")
    n = next(x for x in z.namelist() if x.endswith("Point.csv"))
    testo = z.read(n).decode("cp1252")
    righe = []
    for r in csv.DictReader(io.StringIO(testo), delimiter="\t"):
        if r["TIPO_LOC"] not in ("1", "2"):  # 1 centro abitato, 2 nucleo; non località produttive e case sparse
            continue
        nome = r["NOME"].strip()
        if not nome or nome.lower() in nomi_comuni or not re.search(r"[A-Za-zÀ-ÿ]{3}", nome):
            continue
        sigla = sigla_di.get(int(r["PRO_COM"]), "")
        righe.append(f"{nome};{sigla}")
    return righe


_SUFFISSI = {"street", "road", "avenue", "lane", "drive", "close", "crescent", "gardens", "grove", "terrace", "way", "place",
             "square", "mews", "boulevard", "court", "walk", "row", "hill", "park", "rise", "view", "green", "parade",
             "gate", "mall", "wynd", "vennel", "brae", "loan", "path", "end", "side", "bank", "approach", "yard", "lea"}


def open_names(cart: Path) -> tuple[list, list]:
    z = zipfile.ZipFile(cart / "opname_csv_gb.zip")
    campi = z.read("Doc/OS_Open_Names_Header.csv").decode("utf-8-sig").strip().split(",")
    luoghi, vie = [], Counter()
    for n in z.namelist():
        if not (n.startswith("Data/") and n.endswith(".csv")):
            continue
        for r in csv.reader(io.StringIO(z.read(n).decode("utf-8-sig"))):
            d = dict(zip(campi, r))
            tipo = d.get("LOCAL_TYPE", "")
            if tipo in ("City", "Town", "Village", "Hamlet", "Other Settlement", "Suburban Area"):
                zona = d.get("COUNTY_UNITARY") or d.get("DISTRICT_BOROUGH") or d.get("COUNTRY", "")
                for nome in (d.get("NAME1"), d.get("NAME2")):
                    if nome:
                        luoghi.append(f"{nome};{zona}")
            elif tipo == "Named Road":
                nome = d.get("NAME1", "")
                ultima = nome.split()[-1].lower() if nome else ""
                if nome and ultima not in _SUFFISSI:  # solo le vie SENZA suffisso («Deansgate», «Heol y Felin»)
                    vie[nome] += 1
    return luoghi, sorted(vie)


def census(cart: Path, minimo: int = 200) -> list:
    z = zipfile.ZipFile(cart / "census_2010_surnames.zip")
    out = []
    for r in csv.DictReader(io.StringIO(z.read("Names_2010Census.csv").decode("utf-8"))):
        if r["name"] == "ALL OTHER NAMES" or not r["count"].isdigit() or int(r["count"]) < minimo:
            continue
        out.append(r["name"].title().replace("'S", "'s"))
    return out


def ons(cart: Path, minimo: int = 100) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(cart / "babynames1996to2025.xlsx", read_only=True)
    tot = defaultdict(lambda: [0, 0])
    for foglio, g in (("Table_1", 0), ("Table_2", 1)):
        intest = None
        for r in wb[foglio].iter_rows(values_only=True):
            if r and r[0] == "Name":
                intest = r
                continue
            if intest is None or not r or not r[0]:
                continue
            for k, v in zip(intest, r):
                if isinstance(k, str) and k.endswith("Count") and isinstance(v, int):
                    tot[str(r[0]).strip()][g] += v
    out = []
    for nome, (f, m) in tot.items():
        if f + m >= minimo and re.fullmatch(r"[A-Za-zÀ-ÿ'’-]{2,}", nome):
            out.append(f"{nome};{'F' if f > 2 * m else 'M' if m > 2 * f else 'U'}")
    return out


def wikidata(cart: Path, nome: str) -> list:
    righe = (cart / nome).read_text(encoding="utf-8").splitlines()[1:]
    return [r.strip().strip('"') for r in righe if r.strip() and re.fullmatch(r"[^\d\W][\w'’ -]*", r.strip().strip('"'))]


def main():
    cart = Path(sys.argv[1])
    for f, h in SORGENTI.items():
        reale = sha(cart / f)
        if reale != h:
            sys.exit(f"impronta diversa per {f}: {reale} (attesa {h}). Fermo.")
    print("impronte dei 7 file scaricati: tutte corrette")
    prodotti = []
    rc, sigla_di = comuni(cart)
    prodotti.append(scrivi("istat_comuni.txt", rc))
    nomi_comuni = {r.split(";")[0].lower() for r in rc}
    prodotti.append(scrivi("istat_localita.txt", localita(cart, sigla_di, nomi_comuni)))
    luoghi, vie = open_names(cart)
    prodotti.append(scrivi("osn_luoghi_gb.txt", luoghi))
    prodotti.append(scrivi("osn_vie_gb.txt", vie))
    prodotti.append(scrivi("census_cognomi_us.txt", census(cart)))
    prodotti.append(scrivi("ons_nomi_uk.txt", ons(cart)))
    prodotti.append(scrivi("wikidata_cognomi_it.txt", wikidata(cart, "wikidata_cognomi_it.csv")))
    prodotti.append(scrivi("wikidata_nomi_it.txt", wikidata(cart, "wikidata_nomi_it.csv")))
    fonti = {"sorgenti": {f: {"sha256": h, "fonte": FONTI[f]} for f, h in SORGENTI.items()},
             "prodotti": {p: sha(DIZ / p) for p in prodotti}}
    (DIZ / "FONTI.json").write_text(json.dumps(fonti, ensure_ascii=False, indent=1), encoding="utf-8",
                                   newline="\r\n")  # stessi byte su ogni sistema
    print("FONTI.json scritto")


if __name__ == "__main__":
    main()
